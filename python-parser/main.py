import io
import re
import os
from typing import Dict, List, Optional, Any
from fastapi import FastAPI, File, UploadFile, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

app = FastAPI(
    title="Placement Co-pilot Resume Parser",
    description="Lightweight local PDF & DOCX resume text & structure extraction microservice",
    version="1.0.0",
)

# Enable CORS for Next.js frontend
app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:3000",
        "http://127.0.0.1:3000",
        "http://localhost:3001",
        "http://127.0.0.1:3001",
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# ─── Pydantic Models ─────────────────────────────────────────────────────────

class ContactInfo(BaseModel):
    email: Optional[str] = None
    phone: Optional[str] = None
    links: List[str] = []

class ResumeMeta(BaseModel):
    pageCount: int = 1
    wordCount: int = 0
    bulletCount: int = 0
    hasQuantifiedMetrics: bool = False

class ParseResumeResponse(BaseModel):
    rawText: str
    contact: ContactInfo
    detectedSections: Dict[str, str]
    meta: ResumeMeta

# ─── Regex Helpers ────────────────────────────────────────────────────────────

EMAIL_REGEX = re.compile(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+")
PHONE_REGEX = re.compile(
    r"(?:(?:\+91|91|0)?[\s-]?)?[6-9]\d{9}|(?:\+?1[\s-]?)?\(?\d{3}\)?[\s-]?\d{3}[\s-]?\d{4}"
)
URL_REGEX = re.compile(
    r"https?://(?:www\.)?[-a-zA-Z0-9@:%._+~#=]{1,256}\.[a-zA-Z0-9()]{1,6}\b(?:[-a-zA-Z0-9()@:%_+.~#?&/=]*)|(?:github\.com|linkedin\.com/in|gitlab\.com)/[a-zA-Z0-9_.-]+"
)
METRIC_REGEX = re.compile(r"(\d+(?:\.\d+)?%|\$\d+(?:,\d+)*(?:\.\d+)?k?|₹\d+(?:,\d+)*(?:\.\d+)?(?:L|k)?|\b\d+x\b|\b\d+\+\b)")
BULLET_REGEX = re.compile(r"^\s*([•\-\*▪\u2022\u2023\u25E6\u2043\u2219]|(?:\d+\.))\s+", re.MULTILINE)

SECTION_PATTERNS = {
    "education": re.compile(r"^(?:education|academics|academic background|qualifications)\b", re.IGNORECASE),
    "experience": re.compile(r"^(?:experience|work experience|employment history|professional experience|internships)\b", re.IGNORECASE),
    "projects": re.compile(r"^(?:projects|personal projects|key projects|academic projects|portfolio)\b", re.IGNORECASE),
    "skills": re.compile(r"^(?:skills|technical skills|skills & expertise|core competencies|technologies|tech stack)\b", re.IGNORECASE),
    "achievements": re.compile(r"^(?:achievements|awards|honors|certifications|extracurricular activities|leadership)\b", re.IGNORECASE),
}

# ─── Parser Functions ─────────────────────────────────────────────────────────

def extract_pdf_pymupdf(file_bytes: bytes) -> tuple[str, List[str], int]:
    """Extract text and hyperlinks using PyMuPDF (fitz) preserving layout."""
    try:
        import fitz  # PyMuPDF
        doc = fitz.open(stream=file_bytes, filetype="pdf")
        page_count = len(doc)
        text_blocks = []
        links: List[str] = []

        for page in doc:
            for link in page.get_links():
                uri = link.get("uri")
                if uri and (uri.startswith("http://") or uri.startswith("https://") or "github" in uri or "linkedin" in uri):
                    links.append(uri)

            blocks = page.get_text("blocks")
            blocks.sort(key=lambda b: (b[1], b[0]))
            for b in blocks:
                b_text = b[4].strip()
                if b_text:
                    text_blocks.append(b_text)

        full_text = "\n\n".join(text_blocks)
        return full_text, list(set(links)), page_count
    except Exception as e:
        print(f"[PyMuPDF Notice]: {e}")
        raise

def extract_pdf_pypdf(file_bytes: bytes) -> tuple[str, List[str], int]:
    """Pure-Python PDF text extraction using pypdf."""
    try:
        import pypdf
        reader = pypdf.PdfReader(io.BytesIO(file_bytes))
        pages_text = []
        links: List[str] = []
        for page in reader.pages:
            t = page.extract_text() or ""
            if t.strip():
                pages_text.append(t.strip())
        return "\n\n".join(pages_text), links, len(reader.pages)
    except Exception as e:
        print(f"[pypdf Notice]: {e}")
        return "", [], 1

def extract_pdf_pdfplumber(file_bytes: bytes) -> tuple[str, int]:
    """Fallback PDF extraction using pdfplumber for complex table/multi-column layouts."""
    try:
        import pdfplumber
        with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
            pages_text = []
            for page in pdf.pages:
                t = page.extract_text(layout=True) or page.extract_text() or ""
                if t.strip():
                    pages_text.append(t.strip())
            return "\n\n".join(pages_text), len(pdf.pages)
    except Exception:
        return "", 1

def extract_docx(file_bytes: bytes) -> str:
    """Extract paragraphs and tables from a .docx document."""
    from docx import Document
    doc = Document(io.BytesIO(file_bytes))
    content = []
    for p in doc.paragraphs:
        if p.text.strip():
            content.append(p.text.strip())
    for table in doc.tables:
        for row in table.rows:
            row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                content.append(row_text)
    return "\n\n".join(content)

def split_sections_heuristically(text: str) -> Dict[str, str]:
    """Split clean resume text into standard resume sections."""
    lines = text.split("\n")
    sections: Dict[str, List[str]] = {
        "education": [],
        "experience": [],
        "projects": [],
        "skills": [],
        "achievements": [],
        "other": [],
    }

    current_section = "other"

    for line in lines:
        cleaned_line = line.strip()
        if not cleaned_line:
            continue

        matched_section = None
        # Headers are usually short (< 40 characters) and match our keywords
        if len(cleaned_line) < 40:
            clean_for_match = re.sub(r"[:\-_|]", "", cleaned_line).strip()
            for sec_name, pattern in SECTION_PATTERNS.items():
                if pattern.match(clean_for_match):
                    matched_section = sec_name
                    break

        if matched_section:
            current_section = matched_section
        else:
            sections[current_section].append(cleaned_line)

    return {
        sec: "\n".join(lines).strip()
        for sec, lines in sections.items()
        if sec != "other" and lines
    }

# ─── Endpoints ────────────────────────────────────────────────────────────────

@app.get("/health")
def health_check():
    return {
        "status": "ok",
        "parser": "PyMuPDF + python-docx",
        "service": "placement-copilot-document-parser",
    }

@app.post("/api/parse-resume", response_model=ParseResumeResponse)
async def parse_resume(file: UploadFile = File(...)):
    filename = file.filename or "unknown.pdf"
    content_type = file.content_type or ""
    file_bytes = await file.read()

    # Max 5MB check
    if len(file_bytes) > 5 * 1024 * 1024:
        raise HTTPException(status_code=400, detail="File size exceeds the 5MB limit.")

    raw_text = ""
    links: List[str] = []
    page_count = 1

    ext = os.path.splitext(filename)[1].lower()

    if ext == ".pdf" or "pdf" in content_type:
        try:
            raw_text, links, page_count = extract_pdf_pymupdf(file_bytes)
        except Exception as e:
            print(f"[Notice] PyMuPDF extract failed ({e}). Falling back to pypdf...")
            raw_text, links, page_count = extract_pdf_pypdf(file_bytes)
            if not raw_text.strip():
                raw_text, page_count = extract_pdf_pdfplumber(file_bytes)

        # If primary extractor produced very little text, attempt secondary extractors
        if len(raw_text.strip()) < 50:
            try:
                fallback_text, p_links, p_count = extract_pdf_pypdf(file_bytes)
                if len(fallback_text.strip()) > len(raw_text.strip()):
                    raw_text = fallback_text
                    links = list(set(links + p_links))
                    page_count = p_count
            except Exception:
                pass

    elif ext == ".docx" or "word" in content_type:
        try:
            raw_text = extract_docx(file_bytes)
            page_count = max(1, len(raw_text.split("\n\n")) // 15)
        except Exception as e:
            raise HTTPException(status_code=422, detail=f"Failed to parse DOCX document: {str(e)}")

    elif ext in [".txt", ".text"] or "text/plain" in content_type:
        try:
            raw_text = file_bytes.decode("utf-8")
        except UnicodeDecodeError:
            raw_text = file_bytes.decode("latin-1", errors="ignore")
    else:
        raise HTTPException(
            status_code=400,
            detail=f"Unsupported file format '{ext}'. Please upload a .pdf, .docx, or .txt file.",
        )

    clean_text = raw_text.strip()
    if not clean_text:
        # Fallback raw byte ASCII string extraction (e.g. for flattened or non-standard PDFs)
        chunks = re.findall(r"[\x20-\x7E\t\r\n]{4,}", file_bytes.decode("latin-1", errors="ignore"))
        clean_text = " ".join(chunks).strip()

    if not clean_text or len(clean_text) < 15:
        raise HTTPException(status_code=422, detail="Document appears to be empty or contains only non-extractable scanned images.")

    # Extract Contact Information
    emails = EMAIL_REGEX.findall(clean_text)
    phones = PHONE_REGEX.findall(clean_text)
    text_urls = URL_REGEX.findall(clean_text)

    all_links = list(set([u if u.startswith("http") else f"https://{u}" for u in (links + text_urls)]))
    contact = ContactInfo(
        email=emails[0] if emails else None,
        phone=phones[0] if phones else None,
        links=all_links[:5],
    )

    # Heuristic Sections
    sections = split_sections_heuristically(clean_text)

    # Calculate Metadata
    words = clean_text.split()
    word_count = len(words)
    bullet_matches = BULLET_REGEX.findall(clean_text)
    bullet_count = len(bullet_matches)
    has_metrics = bool(METRIC_REGEX.search(clean_text))

    meta = ResumeMeta(
        pageCount=page_count,
        wordCount=word_count,
        bulletCount=bullet_count,
        hasQuantifiedMetrics=has_metrics,
    )

    return ParseResumeResponse(
        rawText=clean_text,
        contact=contact,
        detectedSections=sections,
        meta=meta,
    )

if __name__ == "__main__":
    import uvicorn
    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=True)
